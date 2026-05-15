"""
自动报告生成模块 - 生成评查报告
"""
from typing import Dict, Any
from datetime import datetime
from pathlib import Path
from jinja2 import Environment, FileSystemLoader
from loguru import logger
import json


class ReportGenerator:
    """报告生成器 - 生成评查报告"""
    
    def __init__(self):
        self.template_dir = Path(__file__).parent.parent / 'config' / 'templates'
        self.output_dir = Path(__file__).parent.parent.parent / 'reports'
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        # 初始化 Jinja2 环境
        self.env = Environment(
            loader=FileSystemLoader(str(self.template_dir)),
            autoescape=True
        )
    
    def generate_report(self, review_result: Dict[str, Any], format: str = 'markdown') -> str:
        """
        生成评查报告
        
        Args:
            review_result: 评查结果
            format: 输出格式 (markdown/docx/pdf)
            
        Returns:
            报告文件路径
        """
        try:
            # 准备数据
            report_data = self._prepare_report_data(review_result)
            
            # 生成报告内容
            if format == 'markdown':
                return self._generate_markdown(report_data)
            elif format == 'docx':
                return self._generate_docx(report_data)
            elif format == 'pdf':
                return self._generate_pdf(report_data)
            else:
                raise ValueError(f"不支持的格式: {format}")
                
        except Exception as e:
            logger.error(f"生成报告失败: {e}")
            raise
    
    def _prepare_report_data(self, review_result: Dict[str, Any]) -> Dict[str, Any]:
        """准备报告数据"""
        now = datetime.now()
        
        data = {
            'report_title': '生态环境行政处罚案卷评查报告',
            'report_date': now.strftime('%Y年%m月%d日'),
            'report_time': now.strftime('%H:%M:%S'),
            'case_id': review_result.get('case_id', '未知'),
            'case_number': review_result.get('case_number', '未知'),
            'case_name': review_result.get('case_name', '未知'),
            'case_type': review_result.get('case_type', '未知'),
            'respondent': review_result.get('respondent', '未知'),
            'organizing_unit': review_result.get('organizing_unit', '未知'),
            'violation': review_result.get('violation', '未知'),
            'penalty_decision': review_result.get('penalty_decision', '未知'),
            'review_date': review_result.get('review_date', now.isoformat()),
            'file_pages': review_result.get('file_pages', 0),
            
            # 合法性评查
            'legality': {
                'has_veto': review_result.get('legality_review', {}).get('has_veto', False),
                'score': review_result.get('legality_review', {}).get('legality_score', 0),
                'pass': review_result.get('legality_review', {}).get('legality_pass', False),
                'veto_items': review_result.get('legality_review', {}).get('veto_items', []),
                'findings': review_result.get('legality_review', {}).get('findings', [])
            },
            
            # 规范性评查
            'normative': {
                'score': review_result.get('normative_review', {}).get('normative_score', 0),
                'document_score': review_result.get('normative_review', {}).get('document_score', 0),
                'basic_deduction': review_result.get('normative_review', {}).get('basic_elements_deduction', 0),
                'findings': review_result.get('normative_review', {}).get('findings', [])
            },
            
            # 裁量基准
            'discretion': review_result.get('discretion', {}),
            
            # 证据链
            'evidence': review_result.get('evidence_chain', {}),
            
            # 文书完整性
            'documents': review_result.get('documents', {}),
            
            # 综合结果
            'comprehensive': {
                'score': review_result.get('comprehensive_score', 0),
                'grade': review_result.get('comprehensive_grade', '未知'),
                'pass': review_result.get('comprehensive_pass', False)
            }
        }
        
        return data
    
    def _generate_markdown(self, data: Dict[str, Any]) -> str:
        """生成 Markdown 报告"""
        # 检查模板是否存在
        template_path = self.template_dir / 'review_report.md'
        if not template_path.exists():
            content = self._generate_default_markdown(data)
        else:
            template = self.env.get_template('review_report.md')
            content = template.render(data)
        
        # 保存文件
        filename = f"report_{data['case_id']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md"
        filepath = self.output_dir / filename
        
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(content)
        
        logger.info(f"Markdown 报告已生成: {filepath}")
        return str(filepath)
    
    def _generate_default_markdown(self, data: Dict[str, Any]) -> str:
        """生成默认 Markdown 报告内容"""
        content = f"""# {data['report_title']}

---

## 📋 基本信息

| 项目 | 内容 |
|------|------|
| 报告编号 | {data['case_id']} |
| 案号 | {data['case_number']} |
| 案件名称 | {data['case_name']} |
| 案件类型 | {data['case_type']} |
| 违法当事人 | {data['respondent']} |
| 承办单位 | {data['organizing_unit']} |
| 违法行为 | {data['violation']} |
| 处罚决定 | {data['penalty_decision']} |
| 评查日期 | {data['report_date']} |
| 文件页数 | {data['file_pages']} 页 |

---

## 🎯 评查结论

| 项目 | 得分 | 等级 |
|------|------|------|
| 综合得分 | {data['comprehensive']['score']} 分 | {data['comprehensive']['grade']} |
| 是否通过 | {'✅ 通过' if data['comprehensive']['pass'] else '❌ 不通过'} | |

---

## 🔍 合法性评查

### 一票否决情况
- 是否触发否决：{'❌ 是' if data['legality']['has_veto'] else '✅ 否'}
- 合法性得分：{data['legality']['score']} 分

### 否决项详情
{self._format_veto_items(data['legality']['veto_items'])}

### 发现问题
{self._format_findings(data['legality']['findings'])}

---

## 📊 规范性评分

### 评分结果
| 项目 | 得分 |
|------|------|
| 文书得分 | {data['normative']['document_score']}/100 |
| 基本要素扣分 | -{data['normative']['basic_deduction']} |
| 规范性得分 | {data['normative']['score']} 分 |

### 发现问题
{self._format_findings(data['normative']['findings'])}

---

## ⚖️ 裁量基准审查

### 审查结果
| 项目 | 内容 |
|------|------|
| 是否适用 | {'✅ 是' if data['discretion'].get('applicable') else '❌ 否'} |
| 法定幅度 | {data['discretion'].get('legal_range', '-')} |
| 裁量百分值 | {data['discretion'].get('total_percentage', 0) * 100:.0f}% |
| 计算金额 | {data['discretion'].get('calculated_fine', 0)} 万元 |
| 案卷记载 | {data['discretion'].get('recorded_fine', 0)} 万元 |
| 合理性评估 | {data['discretion'].get('reasonableness', '未知')} |

### 裁量因素
{self._format_discretion_factors(data['discretion'].get('factors', []))}

---

## 📑 证据链分析

### 分析结果
| 项目 | 评估 |
|------|------|
| 证据完整度 | {data['evidence'].get('completeness_rate', 0)}% |
| 合法性 | {data['evidence'].get('legality', '未知')} |
| 关联性 | {data['evidence'].get('relevance', '未知')} |
| 客观性 | {data['evidence'].get('objectivity', '未知')} |

### 存在问题
{self._format_findings(data['evidence'].get('problems', []))}

---

## 📄 文书完整性

### 检查结果
- 完整度：{data['documents'].get('completeness_rate', 0)}%
- 缺失文书：{', '.join(data['documents'].get('missing_documents', [])) or '无'}

---

## 💡 改进建议

{self._format_suggestions(data)}

---

## 📝 评查人员

> 生成时间：{data['report_date']} {data['report_time']}
> 系统版本：v1.0.0

---
"""
        return content
    
    def _generate_docx(self, data: Dict[str, Any]) -> str:
        """生成 Word 报告"""
        try:
            from docxtpl import DocxTemplate
            
            template_path = self.template_dir / 'review_report.docx'
            if template_path.exists():
                doc = DocxTemplate(str(template_path))
            else:
                # 创建简单的 Word 文档
                from docx import Document
                doc = Document()
                doc.add_heading(data['report_title'], level=1)
                doc.add_paragraph(f"报告日期：{data['report_date']}")
                doc.add_paragraph(f"案号：{data['case_number']}")
                doc.add_paragraph(f"案件名称：{data['case_name']}")
                doc.add_heading('评查结论', level=2)
                doc.add_paragraph(f"综合得分：{data['comprehensive']['score']} 分")
                doc.add_paragraph(f"等级：{data['comprehensive']['grade']}")
            
            filename = f"report_{data['case_id']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            filepath = self.output_dir / filename
            
            if template_path.exists():
                doc.render(data)
            
            doc.save(str(filepath))
            logger.info(f"Word 报告已生成: {filepath}")
            return str(filepath)
            
        except ImportError:
            logger.warning("未安装 python-docx/docxtpl，将生成 Markdown 报告")
            return self._generate_markdown(data)
    
    def _generate_pdf(self, data: Dict[str, Any]) -> str:
        """生成 PDF 报告"""
        try:
            from weasyprint import HTML
            
            # 先生成 HTML
            html_content = self._generate_html(data)
            
            filename = f"report_{data['case_id']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            filepath = self.output_dir / filename
            
            HTML(string=html_content).write_pdf(str(filepath))
            logger.info(f"PDF 报告已生成: {filepath}")
            return str(filepath)
            
        except ImportError:
            logger.warning("未安装 weasyprint，将生成 Markdown 报告")
            return self._generate_markdown(data)
    
    def _generate_html(self, data: Dict[str, Any]) -> str:
        """生成 HTML 内容"""
        return f"""
<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <title>{data['report_title']}</title>
    <style>
        body {{ font-family: 'Microsoft YaHei', sans-serif; margin: 40px; }}
        h1 {{ color: #1a508b; border-bottom: 2px solid #1a508b; padding-bottom: 10px; }}
        h2 {{ color: #2d5a87; margin-top: 30px; }}
        table {{ border-collapse: collapse; width: 100%; margin: 10px 0; }}
        th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
        th {{ background-color: #f2f2f2; }}
        .pass {{ color: green; font-weight: bold; }}
        .fail {{ color: red; font-weight: bold; }}
        .score {{ font-size: 24px; font-weight: bold; color: #1a508b; }}
    </style>
</head>
<body>
    <h1>{data['report_title']}</h1>
    <p>生成时间：{data['report_date']} {data['report_time']}</p>
    
    <h2>基本信息</h2>
    <table>
        <tr><th>案号</th><td>{data['case_number']}</td></tr>
        <tr><th>案件名称</th><td>{data['case_name']}</td></tr>
        <tr><th>当事人</th><td>{data['respondent']}</td></tr>
        <tr><th>承办单位</th><td>{data['organizing_unit']}</td></tr>
    </table>
    
    <h2>评查结论</h2>
    <p class="score">综合得分：{data['comprehensive']['score']} 分</p>
    <p>等级：{data['comprehensive']['grade']}</p>
    <p>是否通过：<span class="{'pass' if data['comprehensive']['pass'] else 'fail'}">
        {'✅ 通过' if data['comprehensive']['pass'] else '❌ 不通过'}
    </span></p>
</body>
</html>
"""
    
    def _format_veto_items(self, veto_items: list) -> str:
        """格式化否决项"""
        if not veto_items:
            return "- 无"
        
        items = []
        for item in veto_items:
            if isinstance(item, dict):
                items.append(f"- 序号{item.get('number', '?')}: {item.get('description', '')}")
            else:
                items.append(f"- {item}")
        
        return '\n'.join(items)
    
    def _format_findings(self, findings: list) -> str:
        """格式化发现问题"""
        if not findings:
            return "- 无"
        return '\n'.join(f"- {f}" for f in findings)
    
    def _format_discretion_factors(self, factors: list) -> str:
        """格式化裁量因素"""
        if not factors:
            return "- 无"
        
        items = []
        for factor in factors:
            percentage = factor.get('percentage', 0)
            sign = '+' if percentage > 0 else ''
            items.append(f"- {factor.get('factor_name', '')}: {factor.get('case_value', '')} ({sign}{percentage * 100:.0f}%)")
        
        return '\n'.join(items)
    
    def _format_suggestions(self, data: Dict) -> str:
        """格式化改进建议"""
        suggestions = []
        
        # 基于评查结果生成建议
        if data['legality']['has_veto']:
            suggestions.append("1. 针对合法性否决项进行整改")
        
        if data['normative']['score'] < 30:
            suggestions.append("2. 加强文书规范性管理")
        
        if data['evidence'].get('completeness_rate', 0) < 60:
            suggestions.append("3. 补充完善证据材料")
        
        if data['discretion'].get('reasonableness') not in ['合理', '基本合理']:
            suggestions.append("4. 审查裁量基准适用是否准确")
        
        if data['documents'].get('missing_documents'):
            suggestions.append("5. 补充缺失的文书材料")
        
        if not suggestions:
            suggestions.append("1. 案卷评查通过，继续保持")
        
        return '\n'.join(suggestions)


# 全局实例
report_generator = ReportGenerator()
